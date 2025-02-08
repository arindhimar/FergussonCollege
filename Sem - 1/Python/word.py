from docx import Document

# Create a new Document
doc = Document()

# Title
doc.add_heading('JavaScript Function Declarations and Expressions', level=1)

# Section: Difference Between Function Declaration and Function Expression
doc.add_heading('Difference Between Function Declaration and Function Expression', level=2)

# Add content
doc.add_paragraph(
    "In JavaScript, functions can be defined in two primary ways: function declarations and function expressions."
)

# Function Declaration
doc.add_heading('Function Declaration', level=3)
doc.add_paragraph(
    "A function declaration defines a named function. It is hoisted, meaning it can be called before its definition in the code."
)
doc.add_paragraph("Example:")
doc.add_paragraph("function greet(name) { return `Hello, ${name}!`; }")

# Function Expression
doc.add_heading('Function Expression', level=3)
doc.add_paragraph(
    "A function expression defines a function as part of an expression. It can be either named or anonymous."
)
doc.add_paragraph("Example:")
doc.add_paragraph("const greet = function(name) { return `Hello, ${name}!`; };")

# Key Differences
doc.add_heading('Key Differences', level=2)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Feature'
hdr_cells[1].text = 'Function Declaration'
hdr_cells[2].text = 'Function Expression'

data = [
    ("Hoisting", "Hoisted to the top of the scope", "Not hoisted; cannot be called before defined"),
    ("Syntax", "function name() { ... }", "const name = function() { ... }"),
    ("Naming", "Always named", "Can be named or anonymous"),
    ("Scope", "Available in the entire scope", "Available only after the expression is evaluated"),
]

for feature, decl, expr in data:
    row_cells = table.add_row().cells
    row_cells[0].text = feature
    row_cells[1].text = decl
    row_cells[2].text = expr

# Section: JavaScript Program to Add, Multiply, and Subtract Two Numbers
doc.add_heading('JavaScript Program to Add, Multiply, and Subtract Two Numbers', level=2)

# Using Function Declarations
doc.add_heading('Using Function Declarations', level=3)
doc.add_paragraph(
    "function add(a, b) { return a + b; }\n"
    "function subtract(a, b) { return a - b; }\n"
    "function multiply(a, b) { return a * b; }\n"
    "// Example usage\n"
    "const num1 = 10;\n"
    "const num2 = 5;\n"
    "console.log('Using Function Declarations:');\n"
    "console.log(`Addition: ${add(num1, num2)}`);\n"
    "console.log(`Subtraction: ${subtract(num1, num2)}`);\n"
    "console.log(`Multiplication: ${multiply(num1, num2)}`);"
)

# Using Function Expressions
doc.add_heading('Using Function Expressions', level=3)
doc.add_paragraph(
    "const addExpr = function(a, b) { return a + b; };\n"
    "const subtractExpr = function(a, b) { return a - b; };\n"
    "const multiplyExpr = function(a, b) { return a * b; };\n"
    "// Example usage\n"
    "console.log('Using Function Expressions:');\n"
    "console.log(`Addition: ${addExpr(num1, num2)}`);\n"
    "console.log(`Subtraction: ${subtractExpr(num1, num2)}`);\n"
    "console.log(`Multiplication: ${multiplyExpr(num1, num2)}`);"
)

# Save the document
doc.save('JavaScript_Functions.docx')

print("Word document created successfully.")