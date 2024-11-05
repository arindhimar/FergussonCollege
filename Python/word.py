from docx import Document

# Create a new Document
doc = Document()
doc.add_heading('Overview: CORS, HLS, and FFmpeg', level=1)

# CORS Section
doc.add_heading('CORS (Cross-Origin Resource Sharing)', level=2)
doc.add_paragraph(
    "Definition: CORS is a security feature implemented by web browsers to restrict how resources on a web page "
    "are requested from different domains. When a web page makes a request to a different domain (cross-origin request), "
    "the browser checks if the server explicitly allows it."
)
doc.add_paragraph(
    "Use Case: In web applications, CORS allows the server to permit resources (like APIs, images, videos) to be accessed by "
    "web pages hosted on different origins."
)
doc.add_paragraph(
    "Why It’s Important: By default, browsers block cross-origin requests for security. Enabling CORS in our Flask application "
    "(using flask_cors library) allows us to handle these requests safely and ensure the browser doesn't block access to resources."
)

# HLS Section
doc.add_heading('HLS (HTTP Live Streaming)', level=2)
doc.add_paragraph(
    "Definition: HLS (HTTP Live Streaming) is a protocol designed by Apple for streaming audio and video over HTTP. It’s commonly "
    "used for video-on-demand services, live broadcasts, and adaptive streaming."
)
doc.add_heading('How It Works:', level=3)
doc.add_paragraph(
    "1. Segmentation: The original video is broken down into short segments, usually a few seconds each, stored as .ts (Transport Stream) files.\n"
    "2. Playlist File: An .m3u8 file, or playlist, lists these segments and provides URLs to each one, so the player can download and play them sequentially.\n"
    "3. Adaptive Streaming: Different quality levels (bitrates) can be defined in the playlist, allowing the player to adjust playback quality based on network speed."
)
doc.add_heading('Benefits:', level=3)
doc.add_paragraph(
    "- Efficient streaming by loading small segments, leading to less buffering and faster playback.\n"
    "- Adaptive bitrate ensures smooth playback on varying network conditions.\n"
    "- Compatibility with many devices, especially Apple devices, and can be supported on others via libraries like HLS.js."
)

# FFmpeg Section
doc.add_heading('FFmpeg (Fast Forward MPEG)', level=2)
doc.add_paragraph(
    "Definition: FFmpeg is a powerful open-source tool for handling video, audio, and multimedia files. It supports encoding, decoding, "
    "transcoding, streaming, and segmenting media files."
)
doc.add_heading('Use Cases in Our Application:', level=3)
doc.add_paragraph(
    "- Video Processing: FFmpeg is used to convert videos to HLS format, breaking them into segments (.ts files) and generating an .m3u8 playlist file.\n"
    "- Command Usage: In our case, the ffmpeg command is used with parameters to specify input files, codecs, segment duration, playlist type, and output pattern."
)
doc.add_paragraph(
    "Example Command:\n"
    "```bash\n"
    "ffmpeg -i input.mp4 -codec:v libx264 -codec:a aac -hls_time 10 -hls_playlist_type vod "
    "-hls_segment_filename \"segment%03d.ts\" -start_number 0 output/index.m3u8\n"
    "```"
)
doc.add_paragraph(
    "Explanation: This command takes an input video, encodes it using the H.264 codec, breaks it into 10-second .ts segments, "
    "and creates an .m3u8 playlist file."
)

# Save Document
doc.save("POC_Learnings.docx")
